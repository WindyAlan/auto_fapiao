from decimal import Decimal

from openpyxl import Workbook, load_workbook

from verify import compare_fields, reconcile_billing_quantities


def test_compare_fields_all_match():
    """所有字段匹配"""
    excel_row = {
        "invoice_no": "26317000001473818243",
        "invoice_date": "2026/05/19",
        "tax_amount": "3237085.21",
        "total_amount": "28137740.66",
    }
    ocr_fields = {
        "invoice_no": "26317000001473818243",
        "invoice_date": "2026/05/19",
        "tax_amount": "3237085.21",
        "total_amount": "28137740.66",
    }
    diffs = compare_fields(excel_row, ocr_fields, confidence=0.95)
    assert len(diffs) == 0


def test_compare_fields_has_diff():
    """存在差异"""
    excel_row = {
        "invoice_no": "26317000001473818243",
        "invoice_date": "2026/05/19",
        "tax_amount": "3237085.21",
        "total_amount": "28137740.66",
    }
    ocr_fields = {
        "invoice_no": "26317000001473818243",
        "invoice_date": "2026/05/20",
        "tax_amount": "3237085.00",
        "total_amount": "28137740.66",
    }
    diffs = compare_fields(excel_row, ocr_fields, confidence=0.95)
    assert len(diffs) == 2
    assert diffs[0].field_name == "发票日期"
    assert diffs[1].field_name == "税金金额"


def test_reconcile_billing_quantities_updates_all_rows_without_increasing():
    """同一PO多行时，按识别出的逐行数量下调数量。"""
    rows = [
        {"_row_idx": 4, "billing_qty": "3"},
        {"_row_idx": 5, "billing_qty": "4"},
        {"_row_idx": 6, "billing_qty": "3"},
    ]

    updates, error = reconcile_billing_quantities(rows, [Decimal("3"), Decimal("3"), Decimal("3")])

    assert error is None
    assert updates == {5: Decimal("3")}


def test_reconcile_billing_quantities_rejects_an_increase():
    """任何一行需要增加时，不能修改Excel，须报错。"""
    rows = [
        {"_row_idx": 4, "billing_qty": "3"},
        {"_row_idx": 5, "billing_qty": "4"},
        {"_row_idx": 6, "billing_qty": "3"},
    ]

    updates, error = reconcile_billing_quantities(rows, [Decimal("3"), Decimal("5"), Decimal("1")])

    assert updates == {}
    assert "不能增加" in error


def test_reconcile_billing_quantities_reduces_total_when_only_total_is_available():
    """只有汇总数量时，仍只能通过下调各行使总数一致。"""
    rows = [
        {"_row_idx": 4, "billing_qty": "3"},
        {"_row_idx": 5, "billing_qty": "4"},
        {"_row_idx": 6, "billing_qty": "3"},
    ]

    updates, error = reconcile_billing_quantities(rows, [Decimal("9")])

    assert error is None
    assert updates == {5: Decimal("3")}


def test_verify_fills_all_duplicate_po_rows_and_reconciles_quantities(tmp_path, monkeypatch):
    """同一 PO 的每一行都回填发票字段，且数量只能下调。"""
    from verify import verify_invoices

    excel_path = tmp_path / "invoice.xlsx"
    wb = Workbook()
    ws = wb.active
    for row, quantity in enumerate((3, 4, 3), start=4):
        ws.cell(row=row, column=32, value="PO-1")  # AF
        ws.cell(row=row, column=78, value=quantity)  # BZ
    wb.save(excel_path)
    wb.close()

    pdf_dir = tmp_path / "pdfs"
    pdf_dir.mkdir()
    (pdf_dir / "PO-1-123A_invoice.pdf").touch()
    monkeypatch.setattr("verify.extract_pdf_content", lambda _: ("ignored", 1.0))
    monkeypatch.setattr("verify.extract_invoice_fields", lambda _: {
        "invoice_no": "12345678901234567890",
        "invoice_date": "2026/07/19",
        "billing_qty_values": [Decimal("3"), Decimal("3"), Decimal("3")],
    })

    results, output_excel = verify_invoices(str(pdf_dir), str(excel_path))

    output = load_workbook(output_excel).active
    assert [output.cell(row=row, column=78).value for row in range(4, 7)] == [3, 3, 3]
    assert [output.cell(row=row, column=79).value for row in range(4, 7)] == [
        "12345678901234567890",
    ] * 3
    assert [output.cell(row=row, column=80).value for row in range(4, 7)] == ["2026/07/19"] * 3
    assert results[0].diffs[-1].field_name == "本次开票数量"
    assert results[0].diffs[-1].fixed is True


def test_verify_exports_pdf_with_invoice_no_even_when_prior_matching_fails(tmp_path, monkeypatch):
    """未完成前序重命名或Excel匹配的PDF，只要识别到发票号仍应导出。"""
    from verify import verify_invoices

    excel_path = tmp_path / "invoice.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(row=4, column=32, value="PO-1")  # AF
    wb.save(excel_path)
    wb.close()

    pdf_dir = tmp_path / "pdfs"
    pdf_dir.mkdir()
    (pdf_dir / "PO-1-123A_invoice.pdf").touch()
    (pdf_dir / "unrenamed.pdf").touch()
    monkeypatch.setattr("verify.extract_pdf_content", lambda path: (path, 1.0))
    monkeypatch.setattr(
        "verify.extract_invoice_fields",
        lambda text: {"invoice_no": "matched" if "PO-1" in text else "unmatched"},
    )

    verify_invoices(str(pdf_dir), str(excel_path))

    assert sorted(path.name for path in (tmp_path / "pdfs_filled").glob("*.pdf")) == [
        "matched.pdf",
        "unmatched.pdf",
    ]
